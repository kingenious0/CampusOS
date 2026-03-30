import re
import sys
import subprocess

# ── Auto-install python-docx if missing ───────────────────────────────────────
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Installing python-docx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

# ── File paths ────────────────────────────────────────────────────────────────
BRAIN = r"C:\Users\kinge\.gemini\antigravity\brain\35e7fd57-cb25-45dd-bcc6-c4c1c12401b8"
PART1 = BRAIN + r"\report_part1.md"
PART2 = BRAIN + r"\report_part2.md"
OUT   = BRAIN + r"\USTED_NAV_Project_Report.docx"

# ── Read both files ───────────────────────────────────────────────────────────
with open(PART1, encoding="utf-8") as f:
    text1 = f.read()
with open(PART2, encoding="utf-8") as f:
    text2 = f.read()

full_text = text1 + "\n\n" + text2
lines = full_text.split("\n")

# ── Document setup ────────────────────────────────────────────────────────────
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# ── Helpers ───────────────────────────────────────────────────────────────────
def add_heading(doc, text, level):
    p = doc.add_heading(text.strip(), level=level)
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
    p.paragraph_format.space_after  = Pt(6)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    return p

def parse_inline(para, text):
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            para.add_run(text[pos:m.start()])
        whole = m.group(0)
        if whole.startswith('**'):
            r = para.add_run(m.group(2)); r.bold = True
        elif whole.startswith('*'):
            r = para.add_run(m.group(3)); r.italic = True
        elif whole.startswith('`'):
            r = para.add_run(m.group(4))
            r.font.name = 'Courier New'; r.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        para.add_run(text[pos:])

def add_paragraph(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    parse_inline(p, text)
    return p

def flush_table(doc, table_buffer):
    if not table_buffer:
        return
    real_rows = []
    for row in table_buffer:
        cells = [c for c in row.split('|') if c.strip() != '']
        if all(re.match(r'^[-:]+$', c.strip()) for c in cells):
            continue
        real_rows.append(cells)
    if not real_rows:
        return
    col_count = max(len(r) for r in real_rows)
    table = doc.add_table(rows=len(real_rows), cols=col_count)
    table.style = 'Table Grid'
    for i, row_data in enumerate(real_rows):
        for j, cell_text in enumerate(row_data[:col_count]):
            cell = table.rows[i].cells[j]
            cell.text = cell_text.strip()
            if i == 0:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'E8F0FE')
                tcPr.append(shd)
    doc.add_paragraph()

# ── Parse ─────────────────────────────────────────────────────────────────────
i = 0
table_buffer = []
in_code_block = False
code_lines = []

while i < len(lines):
    line = lines[i]

    # Code fence
    if line.strip().startswith('```'):
        if not in_code_block:
            in_code_block = True; code_lines = []
        else:
            in_code_block = False
            if code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after  = Pt(4)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Courier New'; run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
        i += 1; continue

    if in_code_block:
        code_lines.append(line); i += 1; continue

    # Table rows
    if line.startswith('|'):
        table_buffer.append(line); i += 1; continue
    else:
        flush_table(doc, table_buffer); table_buffer = []

    stripped = line.strip()

    # Horizontal rule
    if re.match(r'^-{3,}$', stripped):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '4338CA')
        pBdr.append(bottom); pPr.append(pBdr)
        i += 1; continue

    # Headings
    m = re.match(r'^(#{1,6})\s+(.*)', stripped)
    if m:
        level = min(len(m.group(1)), 4)
        add_heading(doc, m.group(2).strip(), level)
        i += 1; continue

    # Blockquote
    if stripped.startswith('>'):
        content = re.sub(r'^>\s*', '', stripped)
        if not re.match(r'^\[!', content.strip()):
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.left_indent  = Cm(1.0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            run = p.add_run(content)
            run.italic = True
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
        i += 1; continue

    # Unordered list
    m = re.match(r'^[-*+]\s+(.*)', stripped)
    if m:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after  = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        parse_inline(p, m.group(1))
        i += 1; continue

    # Ordered list
    m = re.match(r'^\d+\.\s+(.*)', stripped)
    if m:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after  = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        parse_inline(p, m.group(1))
        i += 1; continue

    # Blank line
    if stripped == '':
        i += 1; continue

    # Normal paragraph
    add_paragraph(doc, stripped)
    i += 1

flush_table(doc, table_buffer)

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(OUT)
print(f"\n✅ Done! Report saved to:\n   {OUT}\n")
